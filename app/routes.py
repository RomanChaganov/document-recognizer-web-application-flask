from app import app, WORK_DIR, TESSDATA_PATH
from app.preprocessing import preprocess
from app.tablepars import pars
from app.textrecognizer import recognize
from app.tablesrecognizer import recognize_generator
from app.pipeline import relation_extract

import base64
import cv2
from io import BytesIO

from flask import render_template
from flask import request
from flask import send_from_directory, send_file

from PIL import Image

from openpyxl import Workbook
import os
import tempfile
from werkzeug.utils import secure_filename

from tesserocr import PyTessBaseAPI, PSM, RIL
from tesserocr import iterate_level


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/favicon.ico')
def favicon():
    return send_from_directory(
        os.path.join(app.root_path, 'static'), 
        'favicon.ico', 
        mimetype='image/vnd.microsoft.icon'
    )


@app.route('/upload', methods=['POST'])
def upload():
    if 'file' not in request.files:
        return redirect(request.url)
    
    file = request.files['file']
    if file.filename == '':
        return redirect(request.url)
    
    delete_stamp = request.form.get('delete_stamp')
    
    if file and allowed_file(file.filename):
        with tempfile.TemporaryDirectory(dir=WORK_DIR) as temp_dir:
            filename = secure_filename(file.filename)
            file.save(os.path.join(temp_dir, filename))
            
            binary_img, image, table_xl, data = processing(filename, temp_dir, delete_stamp == 'on')
            binary_io, image_io = pil_to_byteio(binary_img), pil_to_byteio(image)
            
            binary_base64 = get_base64_from_byteio(binary_io)
            image_base64 = get_base64_from_byteio(image_io)
            table_base64 = get_base64_from_byteio(table_xl)
           
            #return send_file(zip_buffer, mimetype='application/zip', as_attachment=True, download_name='archive.zip')
            
            return render_template('answer.html', data_extracted=image_base64, data_table=table_base64, items=data)

    return '<h2>Not OK</h2>'
    

def pil_to_byteio(image):
    img_buffer = BytesIO()
    image.save(img_buffer, format='JPEG')
    img_buffer.seek(0)
    
    return img_buffer
    

def get_base64_from_byteio(data):
    byte_data = data.getvalue()
    base64_encoded = base64.b64encode(byte_data)
    base64_string = base64_encoded.decode('utf-8')
    
    return base64_string


ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg'}
def allowed_file(filename):
    return '.' in filename and \
        filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def processing(filename, temp_dir, delete_stamp):
    binary, image = preprocess(os.path.join(temp_dir, filename), delete_stamp)

    imgs, struct_sizes = pars(binary)
    workbook = Workbook()
    sheet = None

    with PyTessBaseAPI(path=TESSDATA_PATH, lang='rus+eng') as api:
        words = recognize(imgs[0], api)

        for i, cells_imgs in enumerate(recognize_generator(imgs[1], struct_sizes)):
            if i == 0:
                sheet = workbook.active
            else:
                sheet = workbook.create_sheet()
            
            sheet.title = f'Лист {i + 1}'

            _set_cells(cells_imgs, api, sheet)
        
    #with open(os.path.join(temp_dir, 'text.txt'), 'w', encoding='UTF-8') as f:
    #    f.write(str(words))
    
    binary_img = Image.fromarray(binary)
    image = Image.fromarray(image)
    data = relation_extract(words, image)
    
    table_xl = BytesIO()
    workbook.save(table_xl)
    table_xl.seek(0)
    
    # if sheet is not None:
        #workbook.save(os.path.join(temp_dir, 'tables.xlsx'))
    
    # cv2.imwrite(os.path.join(temp_dir, 'binary.jpg'), binary)
    
    return binary_img, image, table_xl, data


def _set_cells(cells_imgs, api, sheet):
    api.SetPageSegMode(PSM.SINGLE_BLOCK)

    for indexes, cell_img in cells_imgs:
        image = Image.fromarray(cell_img)
        api.SetImage(image)
        api.Recognize()

        iterator = api.GetIterator()
        level = RIL.WORD

        cell_text = []
        for itr in iterate_level(iterator, level):
            try:
                word = itr.GetUTF8Text(level)
            except RuntimeError:
                continue

            if not word.strip():
                continue

            cell_text.append(word)

        text = ' '.join(cell_text)

        sheet.cell(row=indexes[0] + 1, column=indexes[1] + 1, value=text)
